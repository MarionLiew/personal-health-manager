from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook


def extract_csv(path: Path) -> str:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return "\n".join("\t".join(row) for row in csv.reader(stream))


def extract_xlsx(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[str] = []
    for sheet in workbook.worksheets:
        rows.append(f"[{sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            rows.append("\t".join("" if value is None else str(value) for value in row))
    workbook.close()
    return "\n".join(rows)
